from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO
from pathlib import Path

from docx import Document

from app.core.config import get_settings
from app.models.search import SearchSnapshot
from app.schemas.search import ReportQuoteResponse, SearchReportResponse, SegmentReportResponse, TopicBriefResponse

FALLBACK_TEMPLATE_PATHS = [
    Path("/workspace/media-aggregator/Шаблон.docx"),
    Path("/Users/savelii/Downloads/Шаблон.docx"),
]

SOURCE_TYPE_LABELS = {
    "news": "цифровые СМИ",
    "digital_media": "цифровые СМИ",
    "media": "цифровые СМИ",
    "telegram": "Telegram",
    "x": "X",
    "twitter": "X",
    "youtube": "Youtube",
    "tv": "TV",
}

COMPLEMENTARITY_LABELS = {
    "позитивная": "в положительном ключе",
    "критическая": "в критическом ключе",
    "смешанная": "в смешанном ключе",
    "нейтральная": "нейтрально",
    "неясная": "без выраженного единого тона",
}


def build_snapshot_docx(snapshot: SearchSnapshot, report: SearchReportResponse) -> bytes:
    document = Document(str(_resolve_template_path()))
    paragraphs = document.paragraphs

    title_country = _build_title_scope(snapshot)
    event_date = snapshot.created_at.strftime("%d.%m.%y")
    source_types = _build_source_types(snapshot)
    state_segment = _find_segment(report, "state_aligned")
    non_state_segment = _find_segment(report, "critical_non_state") or _fallback_non_state_segment(report)
    top_topics = (non_state_segment.dominant_topics if non_state_segment else [])[:4]
    quotes_by_topic = _group_quotes(report.key_quotes)

    _set_paragraph(paragraphs[0], [(f"Реакция {title_country} на «{snapshot.query_text}»", True)])
    _set_paragraph(paragraphs[1], [(f"({event_date})", False)])
    _set_paragraph(
        paragraphs[3],
        [
            (
                f"В ходе мониторинга по итогам «{snapshot.query_text}» были проанализированы материалы "
                f"{snapshot.total_results} источников – {source_types}.",
                False,
            )
        ],
    )
    _set_paragraph(
        paragraphs[4],
        [
            (
                f"В проправительственных медиа ({state_segment.item_count if state_segment else 0}) "
                f"инфоповод освещался {_format_complementarity(state_segment)}. "
                f"Основные акценты и темы: {_segment_topic_overview(state_segment)}.",
                False,
            )
        ],
    )
    _set_paragraph(
        paragraphs[5],
        [
            (
                f"В оппозиционных, «независимых» или прозападных медиа "
                f"({non_state_segment.item_count if non_state_segment else 0}) наиболее обсуждаемые темы: "
                f"{_segment_topic_names(non_state_segment)}.",
                False,
            )
        ],
    )

    topic_paragraph_indexes = [7, 8, 9, 10]
    for index, paragraph_index in enumerate(topic_paragraph_indexes, start=1):
        topic = top_topics[index - 1] if index - 1 < len(top_topics) else None
        _set_paragraph(
            paragraphs[paragraph_index],
            _topic_line_parts(index, topic, non_state_segment.item_count if non_state_segment else snapshot.total_results),
        )

    _set_paragraph(paragraphs[12], [("Ключевые цитаты:", True)])
    quote_topic_indexes = [13, 16, 19]
    quote_line_indexes = [(14, 15), (17, 18), (20, 21)]
    for group_index, topic_paragraph_index in enumerate(quote_topic_indexes):
        topic = top_topics[group_index] if group_index < len(top_topics) else None
        _set_paragraph(paragraphs[topic_paragraph_index], [(f"Тема {group_index + 1}: {_topic_name(topic)}", False)])
        topic_quotes = quotes_by_topic.get(topic.topic if topic else "", [])[:2]
        for line_index, quote_paragraph_index in enumerate(quote_line_indexes[group_index]):
            quote = topic_quotes[line_index] if line_index < len(topic_quotes) else None
            _set_paragraph(paragraphs[quote_paragraph_index], [(_quote_line(quote), False)])

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output.getvalue()


def _resolve_template_path() -> Path:
    settings = get_settings()
    candidates: list[Path] = []
    if settings.report_docx_template_path:
        candidates.append(Path(settings.report_docx_template_path).expanduser())
    candidates.extend(FALLBACK_TEMPLATE_PATHS)
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError("DOCX template was not found. Set REPORT_DOCX_TEMPLATE_PATH or place Шаблон.docx in Downloads.")


def _set_paragraph(paragraph, parts: list[tuple[str, bool]]) -> None:
    paragraph.clear()
    for text, bold in parts:
        run = paragraph.add_run(text)
        run.bold = bold


def _build_title_scope(snapshot: SearchSnapshot) -> str:
    requested_countries = (snapshot.provider_summary.get("_meta", {}) or {}).get("requested_countries", [])
    if not requested_countries:
        return "мировых СМИ"
    if len(requested_countries) == 1:
        return f"СМИ страны {requested_countries[0]}"
    if len(requested_countries) <= 3:
        return f"СМИ стран {', '.join(requested_countries)}"
    return "СМИ выбранных стран"


def _build_source_types(snapshot: SearchSnapshot) -> str:
    values: list[str] = []
    for item in snapshot.items:
        label = SOURCE_TYPE_LABELS.get((item.source_type or "").lower(), item.source_type)
        if label and label not in values:
            values.append(label)
    return ", ".join(values) if values else "цифровые СМИ"


def _find_segment(report: SearchReportResponse, segment_key: str) -> SegmentReportResponse | None:
    for segment in report.segment_reports:
        if segment.segment == segment_key:
            return segment
    return None


def _fallback_non_state_segment(report: SearchReportResponse) -> SegmentReportResponse | None:
    for segment in report.segment_reports:
        if "прозапад" in segment.label.lower() or "оппозицион" in segment.label.lower():
            return segment
    return None


def _format_complementarity(segment: SegmentReportResponse | None) -> str:
    if segment is None:
        return "без выраженного единого тона"
    return COMPLEMENTARITY_LABELS.get(segment.complementarity.lower(), segment.complementarity.lower())


def _segment_topic_overview(segment: SegmentReportResponse | None) -> str:
    if segment is None or not segment.dominant_topics:
        return "тема не выделилась"
    return "; ".join(_topic_name(topic) for topic in segment.dominant_topics[:3])


def _segment_topic_names(segment: SegmentReportResponse | None) -> str:
    if segment is None or not segment.dominant_topics:
        return "не выделены"
    return ", ".join(_topic_name(topic) for topic in segment.dominant_topics[:4])


def _topic_line_parts(index: int, topic: TopicBriefResponse | None, total_items: int) -> list[tuple[str, bool]]:
    if topic is None:
        return [(f"Тема {index} — не выделена.", False)]
    accents = "; ".join(topic.accents[:2]) if topic.accents else "акценты не выделены"
    summary = _trim_text(topic.summary, 180)
    return [
        (f"Тема {index}", True),
        (f" ({topic.mentions}/{total_items}) – {summary}. Акценты: {accents}.", False),
    ]


def _topic_name(topic: TopicBriefResponse | None) -> str:
    if topic is None:
        return "—"
    return topic.topic


def _group_quotes(quotes: list[ReportQuoteResponse]) -> dict[str, list[ReportQuoteResponse]]:
    grouped: dict[str, list[ReportQuoteResponse]] = defaultdict(list)
    for quote in quotes:
        if len(grouped[quote.topic]) >= 2:
            continue
        grouped[quote.topic].append(quote)
    return grouped


def _quote_line(quote: ReportQuoteResponse | None) -> str:
    if quote is None:
        return "Цитата – —"
    clean_quote = _trim_text(quote.quote.replace("\n", " "), 220)
    return f"Цитата – {clean_quote} — {quote.source_name}"


def _trim_text(text: str, limit: int) -> str:
    normalized = " ".join((text or "").split())
    if len(normalized) <= limit:
        return normalized
    return normalized[: limit - 1].rstrip(" ,;:-") + "…"
