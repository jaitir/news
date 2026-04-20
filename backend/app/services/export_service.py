from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from app.models.search import SearchSnapshot
from app.schemas.search import QueryPlanResponse, SearchReportResponse
from app.services.docx_export import build_snapshot_docx
from app.services.query_planner import build_search_plan


def build_snapshot_workbook(snapshot: SearchSnapshot, report: SearchReportResponse | None = None) -> bytes:
    workbook = Workbook()
    articles_sheet = workbook.active
    articles_sheet.title = "Articles"
    articles_sheet.append(
        [
            "Provider",
            "Source type",
            "Source",
            "Country",
            "Language",
            "Published at",
            "Title",
            "URL",
            "Summary",
            "Pivot summary",
            "Narrative",
            "Emotion",
            "Stance",
            "Alignment",
            "Originality",
            "Weighted score",
            "Curated source",
            "Analysis method",
            "Semantic cluster",
            "Semantic query score",
            "Primary segment",
            "Source profile",
            "Ranking score",
        ]
    )

    for item in snapshot.items:
        analysis = (item.raw_payload or {}).get("_analysis", {})
        articles_sheet.append(
            [
                item.provider,
                item.source_type,
                item.source_name,
                item.source_country,
                item.language,
                item.published_at.isoformat(),
                item.title,
                item.url,
                item.summary,
                analysis.get("pivot_summary", item.summary),
                item.narrative,
                item.emotion,
                item.stance,
                analysis.get("query_alignment", "low"),
                analysis.get("originality", "unknown"),
                analysis.get("weighted_score", item.ranking_score),
                analysis.get("is_curated_source", False),
                analysis.get("analysis_method", "heuristic"),
                analysis.get("cluster_label") or analysis.get("semantic_cluster"),
                analysis.get("semantic_query_score"),
                ((analysis.get("source_profile") or {}).get("primary_segment")),
                ((analysis.get("source_profile") or {}).get("canonical_name")),
                item.ranking_score,
            ]
        )

    narratives_sheet = workbook.create_sheet("Narratives")
    narratives_sheet.append(["Narrative", "Emotion", "Count", "Weighted signal"])
    counters: dict[tuple[str, str], dict[str, float]] = {}
    for item in snapshot.items:
        analysis = (item.raw_payload or {}).get("_analysis", {})
        key = (analysis.get("cluster_label") or item.narrative, item.emotion)
        if key not in counters:
            counters[key] = {"count": 0, "weighted_signal": 0.0}
        counters[key]["count"] += 1
        counters[key]["weighted_signal"] += float(analysis.get("weighted_score", item.ranking_score)) / 100
    for (narrative, emotion), payload in sorted(counters.items(), key=lambda row: (row[1]["weighted_signal"], row[1]["count"]), reverse=True):
        narratives_sheet.append([narrative, emotion, payload["count"], round(payload["weighted_signal"], 2)])

    metadata_sheet = workbook.create_sheet("Snapshot")
    stored_query_plan = (snapshot.provider_summary.get("_meta", {}) or {}).get("query_plan") or {}
    if stored_query_plan:
        query_plan = QueryPlanResponse(**stored_query_plan)
    else:
        base_plan = build_search_plan(snapshot.query_text)
        query_plan = QueryPlanResponse(
            original_query=snapshot.query_text,
            expanded_query=base_plan.expanded_query,
            keyword_query=base_plan.keyword_query,
            anchor_terms=base_plan.anchor_terms,
            entity_aliases=base_plan.entity_aliases,
        )
    metadata_sheet.append(["Query", snapshot.query_text])
    metadata_sheet.append(["Expanded query", query_plan.expanded_query])
    metadata_sheet.append(["Keyword query", query_plan.keyword_query])
    metadata_sheet.append(["Anchor terms", ", ".join(query_plan.anchor_terms)])
    metadata_sheet.append(["Paraphrases", " | ".join(query_plan.paraphrases)])
    for language, queries in sorted(query_plan.multilingual_queries.items()):
        metadata_sheet.append([f"Queries ({language})", " | ".join(queries)])
    metadata_sheet.append(["Analysis mode", snapshot.provider_summary.get("_meta", {}).get("analysis_mode", "heuristic")])
    executive_summary = snapshot.provider_summary.get("_meta", {}).get("executive_summary") or {}
    metadata_sheet.append(["Executive overview", executive_summary.get("overview", "")])
    metadata_sheet.append(["Lookback days", snapshot.lookback_days])
    metadata_sheet.append(["Requested sources", ", ".join(snapshot.requested_sources)])
    metadata_sheet.append(
        ["Requested countries", ", ".join((snapshot.provider_summary.get("_meta", {}) or {}).get("requested_countries", []))]
    )
    metadata_sheet.append(["Created at", snapshot.created_at.isoformat()])
    metadata_sheet.append(["Total results", snapshot.total_results])

    if report:
        report_sheet = workbook.create_sheet("Report Overview")
        report_sheet.append(["Report title", report.report_title])
        report_sheet.append(["Overview", report.overview])
        report_sheet.append(["Classified items", report.registry_coverage.classified_items])
        report_sheet.append(["Total items", report.registry_coverage.total_items])
        report_sheet.append(["Classified sources", report.registry_coverage.classified_sources])
        report_sheet.append(["Total sources", report.registry_coverage.total_sources])
        report_sheet.append(
            [
                "Source type breakdown",
                ", ".join(f"{key}:{value}" for key, value in sorted(report.registry_coverage.source_type_breakdown.items())),
            ]
        )
        report_sheet.append(
            [
                "Segment breakdown",
                ", ".join(f"{key}:{value}" for key, value in sorted(report.registry_coverage.segment_breakdown.items())),
            ]
        )
        report_sheet.append(
            [
                "Country breakdown",
                ", ".join(f"{key}:{value}" for key, value in sorted(report.registry_coverage.country_breakdown.items())),
            ]
        )
        report_sheet.append(["Cross-segment gaps", " | ".join(report.cross_segment_gaps)])
        report_sheet.append(["Uncovered sources", " | ".join(report.registry_coverage.uncovered_sources)])

        segments_sheet = workbook.create_sheet("Report Segments")
        segments_sheet.append(
            [
                "Segment",
                "Label",
                "Item count",
                "Source count",
                "Source examples",
                "Source types",
                "Countries",
                "Complementarity",
                "Overview",
                "Omitted topics",
            ]
        )
        for segment in report.segment_reports:
            segments_sheet.append(
                [
                    segment.segment,
                    segment.label,
                    segment.item_count,
                    segment.source_count,
                    " | ".join(segment.source_examples),
                    ", ".join(segment.source_types),
                    ", ".join(segment.countries),
                    segment.complementarity,
                    segment.overview,
                    " | ".join(segment.omitted_topics),
                ]
            )

        topics_sheet = workbook.create_sheet("Report Topics")
        topics_sheet.append(
            [
                "Segment",
                "Topic",
                "Mentions",
                "Total in segment",
                "Summary",
                "Accents",
                "Positive",
                "Negative",
                "Neutral",
                "Omitted topics",
            ]
        )
        for segment in report.segment_reports:
            for topic in segment.dominant_topics:
                topics_sheet.append(
                    [
                        segment.label,
                        topic.topic,
                        topic.mentions,
                        topic.total_in_segment,
                        topic.summary,
                        " | ".join(topic.accents),
                        topic.positive_count,
                        topic.negative_count,
                        topic.neutral_count,
                        " | ".join(topic.omitted_in_segments),
                    ]
                )

        quotes_sheet = workbook.create_sheet("Report Quotes")
        quotes_sheet.append(
            [
                "Topic",
                "Quote",
                "Source",
                "Source type",
                "Provider",
                "Emotion",
                "Country",
                "Primary segment",
            ]
        )
        for quote in report.key_quotes:
            quotes_sheet.append(
                [
                    quote.topic,
                    quote.quote,
                    quote.source_name,
                    quote.source_type,
                    quote.provider,
                    quote.emotion,
                    quote.source_country,
                    (quote.source_profile or {}).get("primary_segment"),
                ]
            )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()


def build_snapshot_docx_bytes(snapshot: SearchSnapshot, report: SearchReportResponse) -> bytes:
    return build_snapshot_docx(snapshot, report)
